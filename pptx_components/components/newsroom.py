from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from pptx_components.base import Component, _resolve, add_rect, add_text_box
from pptx_components.theme import Theme

_ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}


class NewsroomStrap(Component):
    """Top newsroom strap with accent fill and centered/edge-aligned label."""

    def __init__(
        self,
        text: str = "WORLD DESK | DEVELOPING",
        align: str = "center",
        fill_rgb: tuple[int, int, int] | None = None,
        text_rgb: tuple[int, int, int] = (255, 255, 255),
        font_name: str = "Arial",
    ):
        if align not in _ALIGN:
            raise ValueError("align must be one of 'left', 'center', 'right'")
        self.text = text
        self.align = align
        self.fill_rgb = fill_rgb
        self.text_rgb = text_rgb
        self.font_name = font_name

    @property
    def min_height(self) -> float:
        return 0.28

    def render(
        self,
        slide,
        x: float,
        y: float,
        width: float,
        height: float,
        theme: Theme | None = None,
    ) -> None:
        if not self.text or not self.text.strip():
            return

        t = _resolve(theme)
        strap = add_rect(
            slide,
            x,
            y,
            width,
            height,
            fill_rgb=self.fill_rgb or t.ACCENT,
            radius=0.0,
        )
        tf = strap.text_frame
        tf.clear()
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = _ALIGN[self.align]
        run = p.add_run()
        run.text = self.text
        run.font.name = self.font_name
        run.font.bold = True
        run.font.size = Pt(max(t.CAPTION + 1, 9))
        run.font.color.rgb = RGBColor(*self.text_rgb)


class AttributionFooter(Component):
    """Bottom attribution/source line for editorial slides."""

    def __init__(
        self,
        text: str,
        align: str = "left",
        color_rgb: tuple[int, int, int] | None = None,
        font_name: str = "Arial",
    ):
        if align not in _ALIGN:
            raise ValueError("align must be one of 'left', 'center', 'right'")
        self.text = text
        self.align = align
        self.color_rgb = color_rgb
        self.font_name = font_name

    @property
    def min_height(self) -> float:
        return 0.2

    def render(
        self,
        slide,
        x: float,
        y: float,
        width: float,
        height: float,
        theme: Theme | None = None,
    ) -> None:
        if not self.text or not self.text.strip():
            return

        t = _resolve(theme)
        add_text_box(
            slide,
            x,
            y,
            width,
            height,
            self.text,
            t.CAPTION,
            bold=False,
            color_rgb=self.color_rgb or t.TEXT_MUTED,
            alignment=_ALIGN[self.align],
            font_name=self.font_name,
            word_wrap=True,
        )


class BylineDateline(Component):
    """Compact byline + dateline text row; no-op when both fields are empty."""

    def __init__(
        self,
        byline: str | None = None,
        dateline: str | None = None,
        separator: str = " | ",
        align: str = "left",
        color_rgb: tuple[int, int, int] | None = None,
        font_name: str = "Arial",
    ):
        if align not in _ALIGN:
            raise ValueError("align must be one of 'left', 'center', 'right'")
        self.byline = byline
        self.dateline = dateline
        self.separator = separator
        self.align = align
        self.color_rgb = color_rgb
        self.font_name = font_name

    @property
    def min_height(self) -> float:
        return 0.18

    def _text(self) -> str:
        byline = (self.byline or "").strip()
        dateline = (self.dateline or "").strip()
        if byline and dateline:
            return f"{byline}{self.separator}{dateline}"
        return byline or dateline

    def render(
        self,
        slide,
        x: float,
        y: float,
        width: float,
        height: float,
        theme: Theme | None = None,
    ) -> None:
        text = self._text()
        if not text:
            return

        t = _resolve(theme)
        add_text_box(
            slide,
            x,
            y,
            width,
            height,
            text,
            t.CAPTION,
            bold=False,
            color_rgb=self.color_rgb or t.TEXT_SECONDARY,
            alignment=_ALIGN[self.align],
            font_name=self.font_name,
            word_wrap=False,
        )
