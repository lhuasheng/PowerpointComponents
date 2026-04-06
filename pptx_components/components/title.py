from __future__ import annotations

from pptx.enum.text import PP_ALIGN

from pptx_components.base import (
    Component, _resolve, add_rect, add_text_box, add_accent_bar,
    apply_fill, apply_no_line,
)
from pptx_components.theme import Theme


class TitleBlock(Component):
    """Full-width title with optional subtitle and top accent bar."""

    def __init__(
        self,
        title: str,
        subtitle: str | None = None,
        style_overrides: dict[str, int | str | bool] | None = None,
    ):
        self.title = title
        self.subtitle = subtitle
        self.style_overrides = style_overrides or {}

    @property
    def min_height(self) -> float:
        return 1.4

    def render(self, slide, x: float, y: float, width: float, height: float,
               theme: Theme | None = None) -> None:
        t = _resolve(theme)
        bar_h = 0.07
        o = self.style_overrides
        title_size = int(o.get("title_size", t.DISPLAY))
        subtitle_size = int(o.get("body_size", o.get("caption_size", t.SUBHEADING)))
        title_font = str(o.get("font_name", "Calibri Light"))
        title_bold = bool(o.get("title_bold", True))

        # Top accent bar — full width
        add_rect(slide, x, y, width, bar_h, fill_rgb=t.ACCENT)

        title_y = y + bar_h + t.XS
        if self.subtitle:
            title_h = height * 0.55
            sub_h = height - title_h - bar_h - t.XS
            add_text_box(slide, x, title_y, width, title_h,
                         self.title, title_size, bold=title_bold,
                         color_rgb=t.TEXT_PRIMARY, font_name=title_font)
            add_text_box(slide, x, title_y + title_h, width, sub_h,
                         self.subtitle, subtitle_size, bold=False,
                         color_rgb=t.TEXT_SECONDARY, font_name=title_font)
        else:
            add_text_box(slide, x, title_y, width, height - bar_h - t.XS,
                         self.title, title_size, bold=title_bold,
                         color_rgb=t.TEXT_PRIMARY, font_name=title_font)


class SectionHeader(Component):
    """Horizontal section break with left accent bar and optional right-aligned badge."""

    def __init__(self, text: str, badge_text: str | None = None,
                 badge_width: float | None = None,
                 style_overrides: dict[str, int | str | bool] | None = None):
        self.text = text
        self.badge_text = badge_text
        self.badge_width = badge_width
        self.style_overrides = style_overrides or {}

    @property
    def min_height(self) -> float:
        return 0.6

    def render(self, slide, x: float, y: float, width: float, height: float,
               theme: Theme | None = None) -> None:
        t = _resolve(theme)
        bar_w = 0.05
        pad = t.SM
        o = self.style_overrides
        title_size = int(o.get("title_size", t.HEADING))
        badge_size = int(o.get("caption_size", o.get("body_size", t.CAPTION)))
        title_bold = bool(o.get("title_bold", True))
        override_font = o.get("font_name")
        title_font = str(override_font) if override_font is not None else "Calibri Light"
        badge_font = str(override_font) if override_font is not None else "Calibri"

        # Estimate badge width from text length to avoid clipping long labels.
        badge_w = 0.0
        if self.badge_text:
            if self.badge_width is not None:
                badge_w = self.badge_width
            else:
                est = 0.65 + (0.07 * len(self.badge_text))
                badge_w = max(1.3, min(2.4, est))

        # Left accent bar
        add_accent_bar(slide, x, y, height, t, width=bar_w)

        # Header text
        text_x = x + bar_w + pad
        text_w = width - bar_w - pad - (badge_w + pad if self.badge_text else 0)
        add_text_box(slide, text_x, y, text_w, height,
                     self.text, title_size, bold=title_bold,
                     color_rgb=t.TEXT_PRIMARY, font_name=title_font)

        # Optional badge — pill shape on the right
        if self.badge_text:
            badge_h = min(height * 0.7, 0.35)
            badge_y = y + (height - badge_h) / 2
            badge_x = x + width - badge_w
            pill = add_rect(slide, badge_x, badge_y, badge_w, badge_h,
                            fill_rgb=t.ACCENT, radius=0.5)
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
            tf = pill.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = self.badge_text
            run.font.name = badge_font
            run.font.size = Pt(badge_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
