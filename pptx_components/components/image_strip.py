from __future__ import annotations

from collections.abc import Sequence

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from pptx_components.base import Component, _resolve, add_rect
from pptx_components.components.image_card import ImageCard
from pptx_components.theme import Theme


ImageStripItem = tuple[str, str | None] | tuple[str, str | None, str | None]


class ImageStrip(Component):
    """Render a horizontal strip of image cards."""

    def __init__(
        self,
        items: Sequence[ImageStripItem],
        gap: float = 0.1,
        caption_position: str = "below",
    ):
        if caption_position not in ("below", "overlay"):
            raise ValueError("caption_position must be 'below' or 'overlay'")
        self.items = list(items)
        self.gap = gap
        self.caption_position = caption_position

    @property
    def min_height(self) -> float:
        return 1.35 if self.caption_position == "below" else 1.15

    def _normalize_item(self, item, idx: int) -> tuple[str, str | None, str | None]:
        if not isinstance(item, tuple):
            raise TypeError(f"items[{idx}] must be a tuple like (image_path, caption[, badge_text])")

        if len(item) == 2:
            image_path, caption = item
            return image_path, caption, None
        if len(item) == 3:
            image_path, caption, badge_text = item
            return image_path, caption, badge_text

        raise ValueError(
            f"items[{idx}] must have 2 or 3 values: (image_path, caption[, badge_text])"
        )

    def _draw_overlay_caption(
        self,
        slide,
        x: float,
        y: float,
        width: float,
        height: float,
        caption: str,
        theme: Theme,
    ) -> None:
        bar_h = min(0.22, max(0.18, height * 0.2))
        bar = add_rect(
            slide,
            x + 0.04,
            y + height - bar_h - 0.04,
            max(width - 0.08, 0.1),
            bar_h,
            fill_rgb=theme.ACCENT,
            radius=0.02,
        )
        tf = bar.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = caption
        run.font.name = "Calibri"
        run.font.bold = True
        run.font.size = Pt(theme.CAPTION)
        run.font.color.rgb = RGBColor(255, 255, 255)

    def render(
        self,
        slide,
        x: float,
        y: float,
        width: float,
        height: float,
        theme: Theme | None = None,
    ) -> None:
        t = _resolve(theme)

        if not self.items:
            return

        total_gap = self.gap * (len(self.items) - 1)
        card_w = (width - total_gap) / len(self.items)
        if card_w <= 0:
            raise ValueError("ImageStrip width is too small for item count and gap")

        for idx, raw in enumerate(self.items):
            image_path, caption, badge_text = self._normalize_item(raw, idx)
            item_x = x + idx * (card_w + self.gap)

            card_caption = caption if self.caption_position == "below" else None
            card = ImageCard(
                image_path=image_path,
                caption=card_caption,
                badge_text=badge_text,
                mode="contain",
                border_rgb=None,
            )
            card.render(slide, item_x, y, card_w, height, theme=t)

            if self.caption_position == "overlay" and caption:
                self._draw_overlay_caption(slide, item_x, y, card_w, height, caption, t)
