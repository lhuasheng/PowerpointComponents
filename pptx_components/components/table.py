from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from pptx_components.base import Component, _resolve, add_rect, set_font
from pptx_components.theme import Theme


class DataTable(Component):
    """Tabular data with accent header row, zebra stripes, and proportional columns.

    Args:
        headers: Column header strings.
        rows: List of rows; each row is a list of cell strings.
        weights: Proportional column widths (e.g. [3,1,1]). Defaults to equal.
        zebra: Alternate row background colors.
        accent_header: Use accent color for the header row.
    """

    ROW_H = 0.35  # inches per row

    def __init__(self, headers: list[str], rows: list[list[str]],
                 weights: list[float] | None = None,
                 zebra: bool = True,
                 accent_header: bool = True):
        if len(headers) == 0:
            raise ValueError("DataTable requires at least one header")
        if weights is not None:
            if len(weights) != len(headers):
                raise ValueError(
                    f"weights length ({len(weights)}) must match header count ({len(headers)})"
                )
            if any(w <= 0 for w in weights):
                raise ValueError("weights must be positive numbers")
        self.headers = headers
        self.rows = rows
        self.weights = weights
        self.zebra = zebra
        self.accent_header = accent_header

    @property
    def min_height(self) -> float:
        return (len(self.rows) + 1) * self.ROW_H

    def _col_widths(self, total_width: float) -> list[float]:
        n = len(self.headers)
        if self.weights:
            total_w = sum(self.weights)
            return [total_width * (w / total_w) for w in self.weights]
        return [total_width / n] * n

    def _render_row(self, slide, cells: list[str], col_widths: list[float],
                    row_x: float, row_y: float, row_h: float,
                    bg_rgb: tuple[int, int, int],
                    text_rgb: tuple[int, int, int],
                    bold: bool, t: Theme,
                    alignment=PP_ALIGN.LEFT) -> None:
        cursor_x = row_x
        for cell, col_w in zip(cells, col_widths):
            bg = add_rect(slide, cursor_x, row_y, col_w, row_h, fill_rgb=bg_rgb)

            from pptx.util import Inches
            tf = bg.text_frame
            tf.word_wrap = False
            from pptx.util import Pt
            tf.margin_left = Inches(t.SM)
            tf.margin_top = Inches(t.XS)
            tf.margin_right = Inches(t.XS)
            tf.margin_bottom = Inches(t.XS)

            p = tf.paragraphs[0]
            p.alignment = alignment
            run = p.add_run()
            run.text = str(cell)
            run.font.name = "Calibri"
            run.font.size = Pt(t.BODY)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*text_rgb)

            cursor_x += col_w

    def render(self, slide, x: float, y: float, width: float, height: float,
               theme: Theme | None = None) -> None:
        t = _resolve(theme)
        expected_cols = len(self.headers)
        for i, row in enumerate(self.rows, start=1):
            if len(row) != expected_cols:
                raise ValueError(
                    f"Row {i} has {len(row)} columns, expected {expected_cols}. "
                    "All rows must match the header length."
                )
        col_widths = self._col_widths(width)
        cursor_y = y

        # Header row
        header_bg = t.ACCENT if self.accent_header else t.SURFACE_ALT
        header_text = (255, 255, 255) if self.accent_header else t.TEXT_PRIMARY
        self._render_row(slide, self.headers, col_widths,
                         x, cursor_y, self.ROW_H,
                         header_bg, header_text, bold=True, t=t)
        cursor_y += self.ROW_H

        # Data rows
        for i, row in enumerate(self.rows):
            if self.zebra:
                bg = t.SURFACE if i % 2 == 0 else t.SURFACE_ALT
            else:
                bg = t.SURFACE
            self._render_row(slide, row, col_widths,
                             x, cursor_y, self.ROW_H,
                             bg, t.TEXT_PRIMARY, bold=False, t=t)
            cursor_y += self.ROW_H
